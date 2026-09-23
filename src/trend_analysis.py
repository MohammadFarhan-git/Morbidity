"""
Multi-Month Follow-Up Adherence & Resurgery Rate Trend Analysis Module.
Processes 3 consecutive monthly surgical datasets, builds trend comparisons,
and exports executive Excel workbooks and multi-page PDF chart reports.
"""

import io
import os
import re
from datetime import datetime
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages

# ── Clinical Procedure & Prefix Constants ──
CAMPUS_SHEETS = [
    'KAR Campus Data',
    'KVC Campus Data',
    'GMR Campus Data',
    'MTC Campus Data',
    'YSR Campus Data',
]

REPEAT_KP_PROCS = {
    'THERAPUETIC PENETRATING KERATOPLASTY (TH PK)',
    'INTRAOCULAR ANTIBIOTIC INJECTION (IOAB),THERAPUETIC PENETRATING KERATOPLASTY (TH PK)',
    'DESCEMETS STRIPPING AUTOMATED ENDOTHELIAL KERATOPLASTY (DSAEK)',
    'ANTERIOR VITRECTOMY,THERAPUETIC PENETRATING KERATOPLASTY (TH PK)',
    'PENETRATING KERATOPLASTY (PK)',
    'DESCEMET MEMBRANE ENDOTHLIAL KERATOPLASTY (DMEK)',
    'ECCE + IOL,PENETRATING KERATOPLASTY (PK),TARSORRAPHY',
}

PREFIX_MAP = {"PN": "P", "NP": "N", "NPC": "CC", "PNC": "CC", "NPN": "N", "PNP": "P"}


def extract_month_label(filename_or_str: str, fallback_idx: int = 1) -> str:
    """Extract month name and year (if present) from a filename string."""
    if not filename_or_str:
        return f"Month {fallback_idx}"
    months = {
        'jan': 'January', 'feb': 'February', 'mar': 'March', 'apr': 'April',
        'may': 'May', 'jun': 'June', 'jul': 'July', 'aug': 'August',
        'sep': 'September', 'oct': 'October', 'nov': 'November', 'dec': 'December'
    }
    s = str(filename_or_str).lower()
    for m_short, m_full in months.items():
        match = re.search(rf'({m_short}[a-z]*)\s*[-_]?\s*(\d{{4}})?', s)
        if match:
            year = match.group(2)
            return f"{m_full} {year}" if year else m_full
    return f"Month {fallback_idx}"


def correct_mrno(mrno) -> str:
    if pd.isna(mrno):
        return ""
    mrno_str = str(mrno).strip()
    match = re.search(r"(?:.*-)?([A-Z]+)(\d+)$", mrno_str)
    if match:
        prefix, digits = match.groups()
        new_prefix = PREFIX_MAP.get(prefix, prefix)
        return re.sub(r"([A-Z]+)(\d+)$", new_prefix + digits, mrno_str)
    return mrno_str


def categorize_repeat_surgery(x) -> str:
    if pd.isna(x):
        return "Others"
    x_str = str(x).strip().upper()
    if x_str == "REBUBBLING":
        return "REBUBBLING"
    if x_str == "WOUND RESUTURING":
        return "WOUND_RESUTURING"
    if x_str in REPEAT_KP_PROCS:
        return "KP"
    return "Others"


def load_and_clean_month_df(file_source) -> pd.DataFrame:
    """Load campus sheets from an excel file buffer or path, standardizing dates and columns."""
    dfs = []
    # If file_source is bytes / UploadedFile / path:
    try:
        xl = pd.ExcelFile(file_source)
        available = xl.sheet_names
        for s in CAMPUS_SHEETS:
            if s in available:
                d = pd.read_excel(xl, sheet_name=s, header=4)
                dfs.append(d)
        if not dfs:
            for s in available:
                d = pd.read_excel(xl, sheet_name=s, header=4)
                dfs.append(d)
    except Exception:
        # Fallback to direct read
        df = pd.read_excel(file_source)
        dfs.append(df)

    if not dfs:
        return pd.DataFrame()

    df = pd.concat(dfs, ignore_index=True)

    # Clean object columns
    for col in df.select_dtypes(include=["object", "string"]).columns:
        df[col] = df[col].apply(lambda x: x.strip().upper() if isinstance(x, str) else x)

    # Date parsing
    for col in ["visit_date", "surg_date", "dob", "fup_surg_date"]:
        if col in df.columns:
            parsed = pd.to_datetime(df[col], errors='coerce')
            numeric = pd.to_numeric(df[col], errors='coerce')
            valid_numeric = numeric.where(numeric.between(1, 100000))
            excel_dates = pd.Timestamp("1899-12-30") + pd.to_timedelta(valid_numeric, unit="D", errors='coerce')
            df[col] = parsed.fillna(excel_dates)

    if "tp_mrno" in df.columns:
        df["tp_mrno_corrected"] = df["tp_mrno"].apply(correct_mrno)
    else:
        df["tp_mrno_corrected"] = ""

    for required_col in ["surg_eye", "surg_proc_group"]:
        if required_col not in df.columns:
            df[required_col] = "UNKNOWN"

    df["id"] = (
        df["tp_mrno_corrected"].astype(str) + "_" +
        df["surg_eye"].astype(str) + "_" +
        df["surg_proc_group"].astype(str) + "_" +
        df["surg_date"].dt.strftime("%Y-%m-%d").fillna("")
    )

    df = df.dropna(how="all").dropna(subset=["id"]).sort_values(["id", "surg_date", "visit_date"])
    df["days_after_surgery"] = (df["visit_date"] - df["surg_date"]).dt.days
    return df


def process_single_month(df: pd.DataFrame) -> dict:
    """Extract primary surgeries, 1M adherence, and 1M repeat surgeries for a single month."""
    for col_req in ['advise_surg', 'fup_surg_done', 'fup_done_surg_proc']:
        if col_req not in df.columns:
            df[col_req] = None

    # Identify repeat keratoplasties
    repeat_mask = (
        (df['visit_date'] > df['surg_date']) &
        (df['advise_surg'] == 'YES') &
        (df['fup_surg_done'] == 'YES') &
        (df['fup_done_surg_proc'].isin(REPEAT_KP_PROCS))
    )
    repeat_dates = (
        df.loc[repeat_mask]
        .groupby('id', as_index=False)['fup_surg_date'].min()
        .rename(columns={'fup_surg_date': 'first_repeat_keratoplasty_date'})
    )
    visit_df = df.merge(repeat_dates, on='id', how='left')
    visit_df_primary = visit_df[
        visit_df['first_repeat_keratoplasty_date'].isna() |
        (visit_df['visit_date'] < visit_df['first_repeat_keratoplasty_date'])
    ].copy()

    # One row per primary surgery
    surgery_df = visit_df_primary.groupby('id', as_index=False).first()

    # 1-Month Follow-Up Adherence (Strictly 11 <= days_after_surgery <= 45)
    fup_1m = visit_df_primary[visit_df_primary['days_after_surgery'].between(11, 45)]
    adherent_ids = set(fup_1m['id'].unique())
    surgery_df['is_adherent_1m'] = surgery_df['id'].isin(adherent_ids)

    # 1-Month Resurgeries (Strictly 11 <= days_after_surgery <= 45, advise_surg == YES, fup_surg_done == YES)
    repeat_1m = visit_df_primary[
        (visit_df_primary['days_after_surgery'].between(11, 45)) &
        (visit_df_primary['advise_surg'] == 'YES') &
        (visit_df_primary['fup_surg_done'] == 'YES')
    ].copy()
    repeat_1m['Repeat_Category'] = repeat_1m['fup_done_surg_proc'].apply(categorize_repeat_surgery)

    return {
        "surgery_df": surgery_df,
        "visit_df_primary": visit_df_primary,
        "repeat_1m": repeat_1m,
    }


def build_3month_trend_tables(month_results: dict[str, dict]) -> dict[str, pd.DataFrame]:
    """
    Given dict of {month_label: process_single_month_output},
    build side-by-side trend tables for Adherence and Resurgery (Overall, Campus, Surgery Type).
    """
    month_keys = list(month_results.keys())
    m1, m2, m3 = month_keys[0], month_keys[1], month_keys[2]

    # Combine all primary surgeries
    all_surgeries = []
    all_repeats = []
    for m in month_keys:
        sdf = month_results[m]["surgery_df"].copy()
        sdf["Month"] = m
        all_surgeries.append(sdf)

        rdf = month_results[m]["repeat_1m"].copy()
        rdf["Month"] = m
        all_repeats.append(rdf)

    comb_surg = pd.concat(all_surgeries, ignore_index=True)
    comb_repeat = pd.concat(all_repeats, ignore_index=True) if all_repeats else pd.DataFrame()

    # ─────────────────────────────────────────────────────────────
    # Table 1: Adherence Trend — Overall
    # ─────────────────────────────────────────────────────────────
    t1_rows = []
    tot_surg = {m: len(month_results[m]["surgery_df"]) for m in month_keys}
    adh_cnt = {m: month_results[m]["surgery_df"]["is_adherent_1m"].sum() for m in month_keys}
    adh_pct = {m: (adh_cnt[m] / tot_surg[m] * 100).round(2) if tot_surg[m] > 0 else 0.0 for m in month_keys}

    pooled_surg = sum(tot_surg.values())
    pooled_adh = sum(adh_cnt.values())
    pooled_pct = round((pooled_adh / pooled_surg * 100), 2) if pooled_surg > 0 else 0.0
    delta_adh = round(adh_pct[m3] - adh_pct[m1], 2)

    t1_rows.append({"Metric": "Total Primary Surgeries", m1: tot_surg[m1], m2: tot_surg[m2], m3: tot_surg[m3], "3-Month Total": pooled_surg, "Trend (M3 - M1)": tot_surg[m3] - tot_surg[m1]})
    t1_rows.append({"Metric": "1M Adherent Surgeries", m1: adh_cnt[m1], m2: adh_cnt[m2], m3: adh_cnt[m3], "3-Month Total": pooled_adh, "Trend (M3 - M1)": adh_cnt[m3] - adh_cnt[m1]})
    t1_rows.append({"Metric": "1M Adherence Rate (%)", m1: f"{adh_pct[m1]}%", m2: f"{adh_pct[m2]}%", m3: f"{adh_pct[m3]}%", "3-Month Total": f"{pooled_pct}%", "Trend (M3 - M1)": f"{delta_adh:+.2f}%"})
    df_adh_overall = pd.DataFrame(t1_rows)

    # ─────────────────────────────────────────────────────────────
    # Helper for Grouped Adherence Trend
    # ─────────────────────────────────────────────────────────────
    def _grouped_adherence_trend(group_col: str, group_label: str) -> pd.DataFrame:
        groups = sorted(comb_surg[group_col].dropna().unique())
        rows = []
        for g in groups:
            r = {group_label: g}
            g_pooled_surg = 0
            g_pooled_adh = 0
            for m in month_keys:
                msurg = month_results[m]["surgery_df"]
                sub = msurg[msurg[group_col] == g]
                cnt = len(sub)
                adh = sub["is_adherent_1m"].sum() if cnt > 0 else 0
                pct = round(adh / cnt * 100, 2) if cnt > 0 else 0.0
                r[f"{m} Surgeries"] = cnt
                r[f"{m} Adherent %"] = pct
                g_pooled_surg += cnt
                g_pooled_adh += adh

            pool_pct = round(g_pooled_adh / g_pooled_surg * 100, 2) if g_pooled_surg > 0 else 0.0
            r["3-Month Surgeries"] = g_pooled_surg
            r["3-Month Adherent %"] = pool_pct
            r["Trend (M3 vs M1)"] = round(r[f"{m3} Adherent %"] - r[f"{m1} Adherent %"], 2)
            rows.append(r)
        return pd.DataFrame(rows)

    df_adh_campus = _grouped_adherence_trend("sap_code", "Campus")
    df_adh_surg = _grouped_adherence_trend("surg_proc_group", "Surgery Procedure")

    # ─────────────────────────────────────────────────────────────
    # Table 4: Resurgery Trend — Overall & Categories
    # ─────────────────────────────────────────────────────────────
    t4_rows = []
    cats = ["Overall Resurgeries", "REBUBBLING", "WOUND_RESUTURING", "KP", "Others"]

    for cat in cats:
        r = {"Repeat Category": "Total Resurgeries" if cat == "Overall Resurgeries" else cat.replace("_", " ")}
        p_rep = 0
        for m in month_keys:
            tot = tot_surg[m]
            reps = month_results[m]["repeat_1m"]
            sub_rep = len(reps) if cat == "Overall Resurgeries" else (reps["Repeat_Category"] == cat).sum()
            pct = round(sub_rep / tot * 100, 2) if tot > 0 else 0.0
            r[f"{m} Count"] = sub_rep
            r[f"{m} %"] = pct
            p_rep += sub_rep

        p_pct = round(p_rep / pooled_surg * 100, 2) if pooled_surg > 0 else 0.0
        r["3-Month Count"] = p_rep
        r["3-Month %"] = p_pct
        r["Trend (M3 vs M1)"] = round(r[f"{m3} %"] - r[f"{m1} %"], 2)
        t4_rows.append(r)

    df_rep_overall = pd.DataFrame(t4_rows)

    # ─────────────────────────────────────────────────────────────
    # Helper for Grouped Resurgery Trend
    # ─────────────────────────────────────────────────────────────
    def _grouped_resurgery_trend(group_col: str, group_label: str) -> pd.DataFrame:
        groups = sorted(comb_surg[group_col].dropna().unique())
        rows = []
        for g in groups:
            r = {group_label: g}
            g_pooled_surg = 0
            g_pooled_rep = 0
            for m in month_keys:
                msurg = month_results[m]["surgery_df"]
                mrep = month_results[m]["repeat_1m"]
                s_cnt = (msurg[group_col] == g).sum()
                r_cnt = (mrep[group_col] == g).sum() if not mrep.empty and group_col in mrep.columns else 0
                pct = round(r_cnt / s_cnt * 100, 2) if s_cnt > 0 else 0.0
                r[f"{m} Surgeries"] = s_cnt
                r[f"{m} Repeat %"] = pct
                g_pooled_surg += s_cnt
                g_pooled_rep += r_cnt

            p_pct = round(g_pooled_rep / g_pooled_surg * 100, 2) if g_pooled_surg > 0 else 0.0
            r["3-Month Surgeries"] = g_pooled_surg
            r["3-Month Repeat %"] = p_pct
            r["Trend (M3 vs M1)"] = round(r[f"{m3} Repeat %"] - r[f"{m1} Repeat %"], 2)
            rows.append(r)
        return pd.DataFrame(rows)

    df_rep_campus = _grouped_resurgery_trend("sap_code", "Campus")
    df_rep_surg = _grouped_resurgery_trend("surg_proc_group", "Surgery Procedure")

    return {
        "1_Trend_Adherence_Overall": df_adh_overall,
        "2_Trend_Adherence_Campus": df_adh_campus,
        "3_Trend_Adherence_SurgType": df_adh_surg,
        "4_Trend_Resurgery_Overall": df_rep_overall,
        "5_Trend_Resurgery_Campus": df_rep_campus,
        "6_Trend_Resurgery_SurgType": df_rep_surg,
    }


def create_trend_excel(trend_tables: dict[str, pd.DataFrame]) -> io.BytesIO:
    """Creates a professionally beautified openpyxl workbook with all 6 trend sheets."""
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    DATA_FONT = Font(name="Segoe UI", size=10, bold=False, color="000000")
    BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="000000")
    TOTAL_ROW_FILL = PatternFill(start_color="EBF1F5", end_color="EBF1F5", fill_type="solid")
    ALT_ROW_FILL = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    WHITE_ROW_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    THIN_SIDE = Side(border_style="thin", color="D9D9D9")
    MEDIUM_BOTTOM = Side(border_style="medium", color="1F4E78")
    DOUBLE_BOTTOM = Side(border_style="double", color="1F4E78")

    DATA_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
    HEADER_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=MEDIUM_BOTTOM)
    TOTAL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=DOUBLE_BOTTOM)

    TAB_COLORS = {
        "1": "2B579A", "2": "27AE60", "3": "D35400",
        "4": "C0392B", "5": "8E44AD", "6": "2980B9"
    }

    for sheet_name, df in trend_tables.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        prefix = sheet_name.split("_")[0]
        if prefix in TAB_COLORS:
            ws.sheet_properties.tabColor = TAB_COLORS[prefix]
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = 'A2'

        # Headers
        headers = list(df.columns)
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = HEADER_BORDER
        ws.row_dimensions[1].height = 28

        # Data rows
        for row_idx, row_data in enumerate(df.itertuples(index=False), 2):
            ws.row_dimensions[row_idx].height = 20
            first_val = str(row_data[0] or "").lower()
            is_total = "total" in first_val or "overall" in first_val
            r_fill = TOTAL_ROW_FILL if is_total else (ALT_ROW_FILL if row_idx % 2 == 0 else WHITE_ROW_FILL)
            r_font = BOLD_FONT if is_total else DATA_FONT
            r_border = TOTAL_BORDER if is_total else DATA_BORDER

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.fill = r_fill
                cell.font = r_font
                cell.border = r_border
                h_name = headers[col_idx - 1]

                if "%" in h_name or "rate" in h_name.lower():
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, (int, float)):
                        cell.number_format = '0.00"%"'
                elif any(k in h_name for k in ["Count", "Surgeries", "Total"]) and "rate" not in h_name.lower():
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, (int, float)):
                        cell.number_format = '#,##0'
                elif "trend" in h_name.lower() or "δ" in h_name.lower():
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, (int, float)):
                        cell.number_format = '+0.00"%";-0.00"%";0.00"%"'
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Column widths
        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(1, ws.max_row + 1))
            h_len = len(headers[col_idx - 1])
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, h_len + 4, 12), 35)

    wb.save(output)
    output.seek(0)
    return output


def create_trend_pdf(trend_tables: dict[str, pd.DataFrame], month_names: list[str]) -> io.BytesIO:
    """
    Renders a multi-page executive PDF report containing high-resolution
    trend charts for Adherence and Resurgeries across M1, M2, M3.
    """
    pdf_buffer = io.BytesIO()
    m1, m2, m3 = month_names[0], month_names[1], month_names[2]

    # Modern color palette
    c_blue = "#1F4E78"
    c_orange = "#D35400"
    c_green = "#27AE60"
    c_purple = "#8E44AD"
    c_red = "#C0392B"
    months = [m1, m2, m3]

    with PdfPages(pdf_buffer) as pdf:
        # ─────────────────────────────────────────────────────────────
        # Page 1: Executive Overview — Adherence & Resurgery Rate Trends
        # ─────────────────────────────────────────────────────────────
        fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(11.69, 8.27), dpi=150)
        fig.suptitle("3-Month Clinical Quality Trends (1-Month Window: 11–45 Days)", fontsize=16, fontweight='bold', color=c_blue, y=0.96)

        df_adh = trend_tables["1_Trend_Adherence_Overall"]
        adh_rates = [
            float(str(df_adh.loc[df_adh["Metric"] == "1M Adherence Rate (%)", m].values[0]).replace('%', ''))
            for m in months
        ]

        ax1.plot(months, adh_rates, marker='o', linewidth=3, markersize=9, color=c_green, label="1M Adherence %")
        for i, val in enumerate(adh_rates):
            ax1.annotate(f"{val:.1f}%", (months[i], val), textcoords="offset points", xytext=(0, 10), ha='center', fontweight='bold', fontsize=11, color=c_green)
        ax1.set_title("1-Month Follow-Up Adherence Rate Trend", fontsize=12, fontweight='bold', pad=12)
        ax1.set_ylabel("Adherence Rate (%)", fontsize=10)
        ax1.set_ylim(0, 105)
        ax1.grid(True, linestyle="--", alpha=0.5)

        df_rep = trend_tables["4_Trend_Resurgery_Overall"]
        rep_rates = [
            float(df_rep.loc[df_rep["Repeat Category"] == "Total Resurgeries", f"{m} %"].values[0])
            for m in months
        ]
        ax2.plot(months, rep_rates, marker='s', linewidth=3, markersize=9, color=c_red, label="1M Resurgery %")
        for i, val in enumerate(rep_rates):
            ax2.annotate(f"{val:.2f}%", (months[i], val), textcoords="offset points", xytext=(0, 10), ha='center', fontweight='bold', fontsize=11, color=c_red)
        ax2.set_title("1-Month Resurgery Rate Trend (11–45 Days)", fontsize=12, fontweight='bold', pad=12)
        ax2.set_ylabel("Resurgery Rate (%)", fontsize=10)
        ax2.set_ylim(0, max(rep_rates + [5]) * 1.35)
        ax2.grid(True, linestyle="--", alpha=0.5)

        plt.tight_layout(rect=[0.05, 0.08, 0.95, 0.92])
        pdf.savefig(fig)
        plt.close(fig)

        # ─────────────────────────────────────────────────────────────
        # Page 2: 1M Follow-up Adherence by Campus
        # ─────────────────────────────────────────────────────────────
        fig, ax = plt.subplots(figsize=(11.69, 8.27), dpi=150)
        df_camp_adh = trend_tables["2_Trend_Adherence_Campus"]
        campuses = df_camp_adh["Campus"].tolist()
        x = np.arange(len(campuses))
        width = 0.25

        for idx, m in enumerate(months):
            rates = df_camp_adh[f"{m} Adherent %"].tolist()
            bars = ax.bar(x + (idx - 1) * width, rates, width, label=m, alpha=0.85)
            for bar in bars:
                h = bar.get_height()
                if h > 0:
                    ax.annotate(f"{h:.1f}%", xy=(bar.get_x() + bar.get_width() / 2, h), xytext=(0, 3), textcoords="offset points", ha='center', va='bottom', fontsize=8, rotation=45)

        ax.set_title("1-Month Follow-Up Adherence Rate by Campus (11–45 Days)", fontsize=14, fontweight='bold', color=c_blue, pad=15)
        ax.set_xticks(x)
        ax.set_xticklabels(campuses, fontsize=11, fontweight='bold')
        ax.set_ylabel("Adherence Rate (%)", fontsize=11)
        ax.set_ylim(0, 115)
        ax.legend(frameon=True, facecolor="#F8F9FA")
        ax.grid(True, axis='y', linestyle="--", alpha=0.5)

        plt.tight_layout(rect=[0.05, 0.08, 0.95, 0.95])
        pdf.savefig(fig)
        plt.close(fig)

        # ─────────────────────────────────────────────────────────────
        # Page 3: 1M Follow-up Adherence by Surgery Type
        # ─────────────────────────────────────────────────────────────
        fig, ax = plt.subplots(figsize=(11.69, 8.27), dpi=150)
        df_surg_adh = trend_tables["3_Trend_Adherence_SurgType"]
        procs = df_surg_adh["Surgery Procedure"].tolist()
        x = np.arange(len(procs))
        width = 0.25

        for idx, m in enumerate(months):
            rates = df_surg_adh[f"{m} Adherent %"].tolist()
            bars = ax.bar(x + (idx - 1) * width, rates, width, label=m, alpha=0.85)
            for bar in bars:
                h = bar.get_height()
                if h > 0:
                    ax.annotate(f"{h:.1f}%", xy=(bar.get_x() + bar.get_width() / 2, h), xytext=(0, 3), textcoords="offset points", ha='center', va='bottom', fontsize=8, rotation=45)

        ax.set_title("1-Month Follow-Up Adherence Rate by Surgery Procedure", fontsize=14, fontweight='bold', color=c_blue, pad=15)
        ax.set_xticks(x)
        ax.set_xticklabels(procs, fontsize=11, fontweight='bold')
        ax.set_ylabel("Adherence Rate (%)", fontsize=11)
        ax.set_ylim(0, 115)
        ax.legend(frameon=True, facecolor="#F8F9FA")
        ax.grid(True, axis='y', linestyle="--", alpha=0.5)

        plt.tight_layout(rect=[0.05, 0.08, 0.95, 0.95])
        pdf.savefig(fig)
        plt.close(fig)

        # ─────────────────────────────────────────────────────────────
        # Page 4: 1M Resurgery Rates by Campus
        # ─────────────────────────────────────────────────────────────
        fig, ax = plt.subplots(figsize=(11.69, 8.27), dpi=150)
        df_camp_rep = trend_tables["5_Trend_Resurgery_Campus"]
        campuses = df_camp_rep["Campus"].tolist()
        x = np.arange(len(campuses))
        width = 0.25

        for idx, m in enumerate(months):
            rates = df_camp_rep[f"{m} Repeat %"].tolist()
            bars = ax.bar(x + (idx - 1) * width, rates, width, label=m, alpha=0.85)
            for bar in bars:
                h = bar.get_height()
                if h > 0:
                    ax.annotate(f"{h:.1f}%", xy=(bar.get_x() + bar.get_width() / 2, h), xytext=(0, 3), textcoords="offset points", ha='center', va='bottom', fontsize=8)

        ax.set_title("1-Month Resurgery Rate by Campus (11–45 Days)", fontsize=14, fontweight='bold', color=c_blue, pad=15)
        ax.set_xticks(x)
        ax.set_xticklabels(campuses, fontsize=11, fontweight='bold')
        ax.set_ylabel("Resurgery Rate (%)", fontsize=11)
        ax.legend(frameon=True, facecolor="#F8F9FA")
        ax.grid(True, axis='y', linestyle="--", alpha=0.5)

        plt.tight_layout(rect=[0.05, 0.08, 0.95, 0.95])
        pdf.savefig(fig)
        plt.close(fig)

        # ─────────────────────────────────────────────────────────────
        # Page 5: 1M Resurgery Rates by Surgery Procedure
        # ─────────────────────────────────────────────────────────────
        fig, ax = plt.subplots(figsize=(11.69, 8.27), dpi=150)
        df_surg_rep = trend_tables["6_Trend_Resurgery_SurgType"]
        procs = df_surg_rep["Surgery Procedure"].tolist()
        x = np.arange(len(procs))
        width = 0.25

        for idx, m in enumerate(months):
            rates = df_surg_rep[f"{m} Repeat %"].tolist()
            bars = ax.bar(x + (idx - 1) * width, rates, width, label=m, alpha=0.85)
            for bar in bars:
                h = bar.get_height()
                if h > 0:
                    ax.annotate(f"{h:.1f}%", xy=(bar.get_x() + bar.get_width() / 2, h), xytext=(0, 3), textcoords="offset points", ha='center', va='bottom', fontsize=8)

        ax.set_title("1-Month Resurgery Rate by Surgery Procedure (11–45 Days)", fontsize=14, fontweight='bold', color=c_blue, pad=15)
        ax.set_xticks(x)
        ax.set_xticklabels(procs, fontsize=11, fontweight='bold')
        ax.set_ylabel("Resurgery Rate (%)", fontsize=11)
        ax.legend(frameon=True, facecolor="#F8F9FA")
        ax.grid(True, axis='y', linestyle="--", alpha=0.5)

        plt.tight_layout(rect=[0.05, 0.08, 0.95, 0.95])
        pdf.savefig(fig)
        plt.close(fig)

        # ─────────────────────────────────────────────────────────────
        # Page 6: 1M Resurgery Category Breakdown Trend
        # ─────────────────────────────────────────────────────────────
        fig, ax = plt.subplots(figsize=(11.69, 8.27), dpi=150)
        df_rep = trend_tables["4_Trend_Resurgery_Overall"]
        categories = ["REBUBBLING", "WOUND RESUTURING", "KP", "Others"]

        x = np.arange(len(categories))
        width = 0.25

        for idx, m in enumerate(months):
            counts = [
                int(df_rep.loc[df_rep["Repeat Category"] == cat, f"{m} Count"].values[0]) if not df_rep.loc[df_rep["Repeat Category"] == cat].empty else 0
                for cat in categories
            ]
            bars = ax.bar(x + (idx - 1) * width, counts, width, label=m, alpha=0.85)
            for bar in bars:
                h = bar.get_height()
                if h > 0:
                    ax.annotate(f"{h}", xy=(bar.get_x() + bar.get_width() / 2, h), xytext=(0, 3), textcoords="offset points", ha='center', va='bottom', fontsize=9, fontweight='bold')

        ax.set_title("1-Month Resurgery Category Counts (11–45 Days)", fontsize=14, fontweight='bold', color=c_blue, pad=15)
        ax.set_xticks(x)
        ax.set_xticklabels(categories, fontsize=11, fontweight='bold')
        ax.set_ylabel("Number of Resurgeries", fontsize=11)
        ax.legend(frameon=True, facecolor="#F8F9FA")
        ax.grid(True, axis='y', linestyle="--", alpha=0.5)

        plt.tight_layout(rect=[0.05, 0.08, 0.95, 0.95])
        pdf.savefig(fig)
        plt.close(fig)

    pdf_buffer.seek(0)
    return pdf_buffer
